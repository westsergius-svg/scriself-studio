from __future__ import annotations

import html
import json
import re
import xml.etree.ElementTree as ET
from collections.abc import Callable, Iterable
from dataclasses import asdict, dataclass, field
from pathlib import Path
from zipfile import ZipFile

from scriborium.core.chapter_text import chapter_text_html_for_editor
from scriborium.core.scriborium_interchange import import_txt as interchange_import_txt
from scriborium.core.paths import app_data_dir
from scriborium.core.scene_fields import default_scene_dict


@dataclass(slots=True)
class ImportOptions:
    split_mode: str = "headings"  # headings | delimiter | single
    delimiter: str = "###"
    create_entity_cards: bool = False


@dataclass(slots=True)
class ImportResult:
    content: dict
    detected_characters: list[str]
    source: Path


@dataclass(slots=True)
class ExportOptions:
    include_preface: bool = True
    include_synopsis: bool = True
    selected_scene_ids: list[str] = field(default_factory=list)
    include_scene_meta: bool = True
    include_plotlines: bool = True
    include_chapter_metadata: bool = True
    include_book_metadata: bool = True
    include_library_characters: bool = True
    include_library_locations: bool = True
    include_library_objects: bool = True
    include_library_plotlines: bool = True
    include_chapter_notes: bool = False
    include_character_arcs: str = "none"  # none | summary | annotations
    include_toc: bool = True
    include_header_footer: bool = False
    include_page_numbers: bool = False
    font_family: str = "Times New Roman"
    font_size: int = 12
    line_spacing: float = 1.3


@dataclass(slots=True)
class ExportTemplate:
    name: str
    fmt: str
    options: ExportOptions

def _chapters_from_interchange(chapters: list[dict]) -> tuple[list[dict], str | None]:
    """Convert chapters produced by interchange import into the editor format.

    Interchange chapters already carry the editor fields; we only need to drop
    the transient ``_book_title`` marker, ensure every scene has the expected
    defaults, and materialize ``chapterTextHtml`` from plain text when missing.
    """
    if not chapters:
        return [], None

    book_title_override: str | None = None
    first = chapters[0]
    raw_title = first.get("_book_title")
    if raw_title:
        stripped = str(raw_title).strip()
        if stripped:
            book_title_override = stripped

    out: list[dict] = []
    for index, chapter in enumerate(chapters, start=1):
        chapter = dict(chapter)
        chapter.pop("_book_title", None)

        scenes = []
        for scene in chapter.get("scenes", []) or []:
            if not isinstance(scene, dict):
                continue
            normalized = default_scene_dict()
            normalized.update(scene)
            scenes.append(normalized)
        chapter["scenes"] = scenes

        if not str(chapter.get("chapterTextHtml", "")).strip():
            chapter["chapterTextHtml"] = chapter_text_html_for_editor(chapter)
        if not chapter.get("id"):
            chapter["id"] = f"chapter_{index:03d}"
        if not chapter.get("number"):
            chapter["number"] = index
        out.append(chapter)
    return out, book_title_override


class ImportService:
    def import_file(self, source: Path, options: ImportOptions) -> ImportResult:
        suffix = source.suffix.lower()
        resolved_source = source
        book_title_override: str | None = None
        if suffix == ".txt":
            text = source.read_text(encoding="utf-8")
            chapters, book_title_override = _chapters_from_interchange(interchange_import_txt(text))
            if not chapters:
                chapters = self._parse_text(text, options)
        elif suffix == ".md":
            text = source.read_text(encoding="utf-8")
            chapters, book_title_override = _chapters_from_interchange(interchange_import_txt(text))
            if not chapters:
                chapters = self._parse_markdown(text, options)
        elif suffix == ".docx":
            paragraphs = self._read_docx_paragraphs(source)
            chapters = self._parse_docx_paragraphs(paragraphs, options)
        else:
            raise ValueError(f"Unsupported import format: {suffix}")

        if not chapters:
            chapters = [self._build_chapter(1, "Глава 1", [self._build_scene(1, "Сцена 1", "")])]

        detected = self._detect_named_entities(chapters)
        library = {"characters": [], "locations": [], "objects": []}
        if options.create_entity_cards:
            library["characters"] = [
                {
                    "id": f"char_{i+1:03d}",
                    "name": name,
                    "description": "",
                    "extra": "",
                    "sceneAppearances": [],
                    "arcSummary": "",
                    "arcPoints": [],
                }
                for i, name in enumerate(detected)
            ]

        book_title = book_title_override or resolved_source.stem
        content = {
            "book": {"title": book_title, "synopsis": "", "preface": ""},
            "chapters": chapters,
            "library": library,
            "timeline": {
                "plotlines": [{"id": "plotline_main", "name": "Главная сюжетная линия", "color": "#4A7DFF"}]
            },
        }
        return ImportResult(content=content, detected_characters=detected, source=resolved_source)

    def _parse_markdown(self, text: str, options: ImportOptions) -> list[dict]:
        lines = text.splitlines()
        chapters: list[dict] = []
        current_chapter_title = "Глава 1"
        current_scene_title = "Сцена 1"
        current_scene_lines: list[str] = []
        current_scenes: list[tuple[str, str]] = []
        chapter_count = 0
        scene_count = 0

        def flush_scene() -> None:
            nonlocal scene_count, current_scene_lines
            body = "\n".join(current_scene_lines).strip()
            if body:
                scene_count += 1
                current_scenes.append((current_scene_title, body))
            current_scene_lines = []

        def flush_chapter() -> None:
            nonlocal chapter_count, current_scenes
            if not current_scenes:
                current_scenes = [("Сцена 1", "")]
            chapter_count += 1
            scenes = [
                self._build_scene(i + 1, title or f"Сцена {i+1}", body)
                for i, (title, body) in enumerate(current_scenes)
            ]
            chapters.append(self._build_chapter(chapter_count, current_chapter_title, scenes))
            current_scenes = []

        for line in lines:
            if line.startswith("# "):
                flush_scene()
                if current_scenes:
                    flush_chapter()
                current_chapter_title = line[2:].strip() or f"Глава {chapter_count+1}"
                current_scene_title = "Сцена 1"
                continue
            if line.startswith("## "):
                flush_scene()
                current_scene_title = line[3:].strip() or f"Сцена {scene_count+1}"
                continue
            current_scene_lines.append(line)

        flush_scene()
        if current_scenes:
            flush_chapter()
        return chapters

    def _parse_text(self, text: str, options: ImportOptions) -> list[dict]:
        if options.split_mode == "single":
            return [self._build_chapter(1, "Глава 1", [self._build_scene(1, "Сцена 1", text.strip())])]
        if options.split_mode == "delimiter":
            delimiter = options.delimiter or "###"
            parts = [p.strip() for p in text.split(delimiter) if p.strip()]
            if not parts:
                parts = [text.strip()]
            scenes = [self._build_scene(i + 1, f"Сцена {i+1}", part) for i, part in enumerate(parts)]
            return [self._build_chapter(1, "Глава 1", scenes)]
        return self._parse_markdown(text, options)

    def _read_docx_paragraphs(self, source: Path) -> list[tuple[str, str]]:
        paragraphs: list[tuple[str, str]] = []
        with ZipFile(source, "r") as archive:
            xml_bytes = archive.read("word/document.xml")
        root = ET.fromstring(xml_bytes)
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}

        for paragraph in root.findall(".//w:p", ns):
            texts = [t.text or "" for t in paragraph.findall(".//w:t", ns)]
            line = "".join(texts).strip()
            if not line:
                continue
            style = ""
            style_node = paragraph.find(".//w:pStyle", ns)
            if style_node is not None:
                style = style_node.attrib.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val", "")
            paragraphs.append((line, style))
        return paragraphs

    def _parse_docx_paragraphs(self, paragraphs: list[tuple[str, str]], options: ImportOptions) -> list[dict]:
        chapters: list[dict] = []
        chapter_title = "Глава 1"
        scene_title = "Сцена 1"
        scene_lines: list[str] = []
        scenes_buf: list[tuple[str, str]] = []
        chapter_idx = 0
        scene_idx = 0

        def flush_scene() -> None:
            nonlocal scene_idx, scene_lines
            body = "\n".join(scene_lines).strip()
            if body:
                scene_idx += 1
                scenes_buf.append((scene_title, body))
            scene_lines = []

        def flush_chapter() -> None:
            nonlocal chapter_idx, scenes_buf
            if not scenes_buf:
                scenes_buf = [("Сцена 1", "")]
            chapter_idx += 1
            scenes = [self._build_scene(i + 1, title, body) for i, (title, body) in enumerate(scenes_buf)]
            chapters.append(self._build_chapter(chapter_idx, chapter_title, scenes))
            scenes_buf = []

        for line, style in paragraphs:
            if style.lower().startswith("heading1"):
                flush_scene()
                if scenes_buf:
                    flush_chapter()
                chapter_title = line
                scene_title = "Сцена 1"
                continue
            if style.lower().startswith("heading2"):
                flush_scene()
                scene_title = line
                continue
            scene_lines.append(line)

        flush_scene()
        if scenes_buf:
            flush_chapter()
        return chapters

    def _detect_named_entities(self, chapters: list[dict]) -> list[str]:
        text = "\n".join(scene.get("text", "") for chapter in chapters for scene in chapter.get("scenes", []))
        pattern = re.compile(r"\b[\u0410-\u042FA-Z][\u0430-\u044Fa-z]{2,}\b")
        found: list[str] = []
        seen: set[str] = set()
        for token in pattern.findall(text):
            if token not in seen:
                seen.add(token)
                found.append(token)
            if len(found) >= 30:
                break
        return found

    def _build_chapter(self, index: int, title: str, scenes: list[dict]) -> dict:
        return {"id": f"chapter_{index:03d}", "title": title, "scenes": scenes}

    def _build_scene(self, index: int, title: str, text: str) -> dict:
        scene = {
            "id": f"scene_{index:03d}",
            "title": title,
            "text": text,
            "characterIds": [],
            "locationIds": [],
            "objectIds": [],
            "plotlineId": "plotline_main",
            "plotlineIds": ["plotline_main"],
            "characterArcPoints": [],
        }
        scene.update(default_scene_dict())
        return scene


def _pdf_wrap(text: str, width: float, font_name: str, size: int) -> list[str]:
    """Wrap a single line of text to the given width in PDF points."""
    try:
        from reportlab.pdfbase.pdfmetrics import stringWidth
    except Exception:
        return [text] if text else []
    words = str(text).split(" ")
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = (current + " " + word).strip() if current else word
        if not current or stringWidth(candidate, font_name, size) <= width:
            current = candidate
        else:
            lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines

class ExportService:
    def __init__(self) -> None:
        self._registry: dict[
            str,
            tuple[str, Callable[[], tuple[bool, str]], Callable[[dict, ExportOptions, Path], Path]],
        ] = {
            "txt": ("TXT", self._always_available, self._export_txt),
            "html": ("HTML", self._always_available, self._export_html),
            "docx": ("DOCX", self._check_docx_available, self._export_docx),
            "pdf": ("PDF", self._check_pdf_available, self._export_pdf),
        }

    def available_formats(self) -> list[tuple[str, str, bool, str]]:
        result: list[tuple[str, str, bool, str]] = []
        for fmt, (label, checker, _) in self._registry.items():
            ok, reason = checker()
            result.append((label, fmt, ok, reason))
        return result

    def render_preview(self, content: dict, fmt: str, options: ExportOptions | None = None) -> str:
        opts = options or ExportOptions()
        if fmt.lower() == "txt":
            text = html.escape(self._to_txt(content, opts))
            return (
                "<html><body><pre style='white-space: pre-wrap; font-family: Consolas, monospace;'>"
                f"{text}</pre></body></html>"
            )
        return self._to_html(content, opts, preview_mode=True, preview_format=fmt.lower())

    def _always_available(self) -> tuple[bool, str]:
        return True, ""

    def _check_docx_available(self) -> tuple[bool, str]:
        if self._has_docx_support():
            return True, ""
        return False, "Requires python-docx"

    def _check_pdf_available(self) -> tuple[bool, str]:
        if self._has_pdf_support():
            return True, ""
        return False, "Requires reportlab"

    def _has_docx_support(self) -> bool:
        try:
            from importlib.util import find_spec
            return find_spec("docx") is not None
        except Exception:
            return False

    def _has_pdf_support(self) -> bool:
        try:
            from importlib.util import find_spec
            return find_spec("reportlab") is not None
        except Exception:
            return False

    def export_document(
        self,
        content: dict,
        fmt: str,
        target: Path,
        options: ExportOptions | None = None,
    ) -> Path:
        fmt_normalized = fmt.lower()
        spec = self._registry.get(fmt_normalized)
        if spec is None:
            raise ValueError(f"Unsupported export format: {fmt}")

        opts = options or ExportOptions()
        target.parent.mkdir(parents=True, exist_ok=True)
        label, checker, adapter = spec
        ok, reason = checker()
        if not ok:
            raise ValueError(f"Export {label} is unavailable: {reason}")
        return adapter(content, opts, target)

    def _export_txt(self, content: dict, options: ExportOptions, target: Path) -> Path:
        target.write_text(self._to_txt(content, options), encoding="utf-8", newline="\n")
        return target

    def _export_html(self, content: dict, options: ExportOptions, target: Path) -> Path:
        target.write_text(self._to_html(content, options), encoding="utf-8", newline="\n")
        return target

    def _export_docx(self, content: dict, options: ExportOptions, target: Path) -> Path:
        return self._to_docx(content, options, target)

    def _export_pdf(self, content: dict, options: ExportOptions, target: Path) -> Path:
        return self._to_pdf(content, options, target)

    def _iter_selected_scenes(self, content: dict, options: ExportOptions) -> Iterable[tuple[dict, list[dict]]]:
        selected = set(options.selected_scene_ids)
        use_all = not selected
        for chapter in content.get("chapters", []):
            chunk = []
            for scene in chapter.get("scenes", []):
                if use_all or scene.get("id") in selected:
                    chunk.append(scene)
            if chunk:
                yield chapter, chunk

    def _volume_blocks(
        self, content: dict, options: ExportOptions
    ) -> Iterable[tuple[dict, list[tuple[dict, list[dict]]]]]:
        """Идёт по томам в их порядке, отдавая (volume, [(chapter, scenes)])."""
        from scriborium.core.volumes import get_volumes

        volumes = get_volumes(content)
        groups: dict[str, list[tuple[dict, list[dict]]]] = {}
        for chapter, scenes in self._iter_selected_scenes(content, options):
            groups.setdefault(chapter.get("volumeId"), []).append((chapter, scenes))
        for volume in volumes:
            items = groups.get(volume["id"])
            if items:
                yield volume, items

    def _volume_outline(
        self, content: dict, options: ExportOptions
    ) -> Iterable[tuple[dict, list[tuple[str, list[str]]]]]:
        """Содержание по томам: (volume, [(chapter_title, [scene_titles])])."""
        for volume, chapter_blocks in self._volume_blocks(content, options):
            entries: list[tuple[str, list[str]]] = []
            for chapter, scenes in chapter_blocks:
                chapter_title = str(chapter.get("title", "Chapter")).strip() or "Chapter"
                scene_titles = [str(scene.get("title", "Scene")).strip() or "Scene" for scene in scenes]
                entries.append((chapter_title, scene_titles))
            if entries:
                yield volume, entries

    @staticmethod
    def _volume_label_text(volume: dict) -> str:
        from scriborium.core.volumes import volume_label

        return volume_label(volume)
    def _entity_name(self, content: dict, kind: str, entity_id: str) -> str:
        for entity in content.get("library", {}).get(kind, []):
            if entity.get("id") == entity_id:
                return entity.get("name", entity_id)
        return entity_id

    def _plotline_name(self, content: dict, plotline_id: str) -> str:
        for plotline in content.get("timeline", {}).get("plotlines", []):
            if plotline.get("id") == plotline_id:
                return plotline.get("name", plotline_id)
        return plotline_id

    def _collect_arc_summary(self, content: dict, options: ExportOptions) -> list[tuple[str, list[tuple[str, int, str]]]]:
        selected_scene_ids = set(options.selected_scene_ids)
        use_all = not selected_scene_ids
        scene_titles: dict[str, str] = {}
        for chapter in content.get("chapters", []):
            for scene in chapter.get("scenes", []):
                scene_id = str(scene.get("id", ""))
                if scene_id and (use_all or scene_id in selected_scene_ids):
                    scene_titles[scene_id] = str(scene.get("title", scene_id))

        summary: list[tuple[str, list[tuple[str, int, str]]]] = []
        for character in content.get("library", {}).get("characters", []):
            points: list[tuple[str, int, str]] = []
            for point in character.get("arcPoints", []):
                scene_id = str(point.get("sceneId", ""))
                if not scene_id or scene_id not in scene_titles:
                    continue
                try:
                    arc_value = int(point.get("arcValue", 0))
                except (TypeError, ValueError):
                    arc_value = 0
                points.append((scene_titles[scene_id], arc_value, str(point.get("description", "")).strip()))
            if points:
                summary.append((str(character.get("name", "Персонаж")), points))
        return summary

    def _scene_arc_annotations(self, content: dict, scene: dict) -> list[str]:
        out: list[str] = []
        for point in scene.get("characterArcPoints", []):
            character_id = str(point.get("characterId", ""))
            if not character_id:
                continue
            character_name = self._entity_name(content, "characters", character_id)
            try:
                arc_value = int(point.get("arcValue", 0))
            except (TypeError, ValueError):
                arc_value = 0
            description = str(point.get("description", "")).strip()
            suffix = f": {description}" if description else ""
            out.append(f"Арка {character_name}: {arc_value:+d}{suffix}")
        return out

    def _scene_meta_lines(self, content: dict, scene: dict, options: ExportOptions) -> list[str]:
        if not options.include_scene_meta:
            return []
        meta: list[str] = []
        if scene.get("dateTime"):
            meta.append(f"Дата/время: {scene.get('dateTime')}")
        if scene.get("goal"):
            meta.append(f"Цель: {scene.get('goal')}")
        if scene.get("conflict"):
            meta.append(f"Конфликт: {scene.get('conflict')}")
        if scene.get("result"):
            meta.append(f"Результат: {scene.get('result')}")
        char_names = [self._entity_name(content, "characters", x) for x in scene.get("characterIds", [])]
        loc_names = [self._entity_name(content, "locations", x) for x in scene.get("locationIds", [])]
        obj_names = [self._entity_name(content, "objects", x) for x in scene.get("objectIds", [])]
        plotline_ids = scene.get("plotlineIds", []) or [scene.get("plotlineId", "plotline_main")]
        plotline_names = [self._plotline_name(content, x) for x in plotline_ids]
        if char_names:
            meta.append("Персонажи: " + ", ".join(char_names))
        if loc_names:
            meta.append("Локации: " + ", ".join(loc_names))
        if obj_names:
            meta.append("Объекты: " + ", ".join(obj_names))
        if options.include_plotlines and plotline_names:
            meta.append("Сюжетные линии: " + ", ".join(plotline_names))
        return meta

    def _scene_heading(self, scene: dict) -> str:
        scene_title = str(scene.get("title", "Сцена")).strip() or "Сцена"
        scene_datetime = str(scene.get("dateTime", "")).strip()
        if scene_datetime:
            return f"{scene_datetime} - {scene_title}"
        return scene_title

    def _chapter_text_plain(self, chapter: dict) -> str:
        plain_text = str(chapter.get("chapterText", "")).strip()
        if plain_text:
            return plain_text
        chapter_html = str(chapter.get("chapterTextHtml", "")).strip()
        if not chapter_html:
            return ""
        # Drop <style>/<script> blocks entirely (their CSS text must not leak into exports).
        text = re.sub(
            r"<\s*(style|script)\b[^>]*>.*?<\s*/\s*(style|script)\s*>",
            "",
            chapter_html,
            flags=re.IGNORECASE | re.DOTALL,
        )
        text = re.sub(r"<\s*br\s*/?\s*>", "\n", text, flags=re.IGNORECASE)
        text = re.sub(r"<\s*li\b[^>]*>", "\n- ", text, flags=re.IGNORECASE)
        text = re.sub(r"</\s*(p|div|h[1-6]|li|tr)\s*>", "\n", text, flags=re.IGNORECASE)
        text = re.sub(r"<[^>]+>", "", text)
        # Collapse 3+ consecutive blank lines.
        text = re.sub(r"[ \t]*\n[ \t]*\n[ \t]*\n+", "\n\n", text)
        return html.unescape(text).strip()

    def _chapter_text_html_fragment(self, chapter: dict) -> str:
        chapter_html = str(chapter.get("chapterTextHtml", "")).strip()
        if chapter_html:
            body_match = re.search(r"<body[^>]*>(.*)</body>", chapter_html, flags=re.IGNORECASE | re.DOTALL)
            fragment = body_match.group(1).strip() if body_match else chapter_html
            if fragment:
                return fragment
        plain_text = self._chapter_text_plain(chapter)
        if not plain_text:
            return ""
        return f"<p>{html.escape(plain_text).replace(chr(10), '<br/>')}</p>"

    def _build_outline(self, content: dict, options: ExportOptions) -> list[tuple[str, list[str]]]:
        outline: list[tuple[str, list[str]]] = []
        for chapter, scenes in self._iter_selected_scenes(content, options):
            chapter_title = str(chapter.get("title", "Chapter")).strip() or "Chapter"
            scene_titles = [str(scene.get("title", "Scene")).strip() or "Scene" for scene in scenes]
            outline.append((chapter_title, scene_titles))
        return outline

    def _chapter_has_scene_text(self, chapter: dict) -> bool:
        return any(str(s.get("text", "")).strip() for s in chapter.get("scenes", []))

    def _chapter_free_text_plain(self, chapter: dict) -> str:
        """Свободный текст главы (после переноса в него текстов сцен)."""
        return self._chapter_text_plain(chapter)

    def _chapter_free_text_html(self, chapter: dict) -> str:
        if self._chapter_has_scene_text(chapter):
            return ""
        return self._chapter_text_html_fragment(chapter)

    def _scene_title_by_id(self, content: dict, scene_id: str) -> str:
        for chapter in content.get("chapters", []):
            for scene in chapter.get("scenes", []):
                if scene.get("id") == scene_id:
                    return str(scene.get("title", scene_id)).strip() or scene_id
        return scene_id

    def _book_metadata_lines(self, book: dict, options: ExportOptions) -> list[str]:
        if not options.include_book_metadata:
            return []
        lines: list[str] = []
        mapping = (
            ("subtitle", "Подзаголовок"),
            ("author", "Автор"),
            ("genre", "Жанр"),
            ("format", "Формат"),
            ("tone", "Тон"),
            ("audience", "Аудитория"),
        )
        for key, label in mapping:
            value = str(book.get(key, "")).strip()
            if value:
                lines.append(f"{label}: {value}")
        return lines

    def _chapter_metadata_lines(self, chapter: dict, options: ExportOptions) -> list[str]:
        if not options.include_chapter_metadata:
            return []
        lines: list[str] = []
        number = chapter.get("number")
        if number not in (None, "", 0):
            lines.append(f"Номер главы: {number}")
        for key, label in (
            ("subtitle", "Подзаголовок"),
            ("epigraph", "Эпиграф"),
            ("synopsis", "Краткое содержание"),
        ):
            value = str(chapter.get(key, "")).strip()
            if value:
                lines.append(f"{label}: {value}")
        if chapter.get("status"):
            lines.append(f"Статус: {chapter.get('status')}")
        target_words = chapter.get("targetWordCount")
        if target_words:
            lines.append(f"Целевой объём (слов): {target_words}")
        return lines

    def _has_additional_materials(self, content: dict, options: ExportOptions) -> bool:
        if self._book_metadata_lines(content.get("book", {}), options):
            return True
        if options.include_character_arcs == "summary" and self._collect_arc_summary(content, options):
            return True
        for chapter, scenes in self._iter_selected_scenes(content, options):
            if self._chapter_metadata_lines(chapter, options):
                return True
            if options.include_chapter_notes and str(chapter.get("notes", "")).strip():
                return True
            for scene in scenes:
                if self._scene_meta_lines(content, scene, options):
                    return True
                if options.include_character_arcs == "annotations" and self._scene_arc_annotations(content, scene):
                    return True
        if self._library_section_enabled(options) and (
            content.get("library") or content.get("timeline")
        ):
            return True
        return False

    def _additional_materials_txt(self, content: dict, options: ExportOptions) -> list[str]:
        body: list[str] = []
        book_meta = self._book_metadata_lines(content.get("book", {}), options)
        if book_meta:
            body.append("Метаданные книги")
            body.append("-" * 16)
            body.extend(book_meta)
            body.append("")
        if options.include_character_arcs == "summary":
            arc_summary = self._collect_arc_summary(content, options)
            if arc_summary:
                body.append("Арки персонажей")
                body.append("-" * 16)
                for character_name, points in arc_summary:
                    body.append(character_name)
                    for scene_title, arc_value, description in points:
                        suffix = f" - {description}" if description else ""
                        body.append(f"  * {scene_title}: {arc_value:+d}{suffix}")
                body.append("")
        for chapter, scenes in self._iter_selected_scenes(content, options):
            chapter_meta = self._chapter_metadata_lines(chapter, options)
            notes = str(chapter.get("notes", "")).strip() if options.include_chapter_notes else ""
            scene_meta: list[tuple[str, list[str]]] = []
            for scene in scenes:
                meta = self._scene_meta_lines(content, scene, options)
                if options.include_character_arcs == "annotations":
                    meta.extend(self._scene_arc_annotations(content, scene))
                if meta:
                    scene_meta.append((self._scene_heading(scene), meta))
            if chapter_meta or notes or scene_meta:
                body.append(str(chapter.get("title", "Глава")))
                body.append("-" * max(5, len(body[-1])))
                body.extend(chapter_meta)
                if notes:
                    body.append(f"Заметки главы: {notes}")
                for heading, meta in scene_meta:
                    body.append(heading)
                    body.extend(f"  {m}" for m in meta)
                body.append("")
        self._append_library_txt(body, content, options)
        return body

    def _additional_materials_html(self, content: dict, options: ExportOptions) -> list[str]:
        body: list[str] = []
        book_meta = self._book_metadata_lines(content.get("book", {}), options)
        if book_meta:
            body.append("<h3>Метаданные книги</h3>")
            body.append("<ul>")
            for item in book_meta:
                label, sep, value = item.partition(": ")
                if sep:
                    body.append(f"<li><b>{html.escape(label)}:</b> {html.escape(value)}</li>")
                else:
                    body.append(f"<li>{html.escape(item)}</li>")
            body.append("</ul>")
        if options.include_character_arcs == "summary":
            arc_summary = self._collect_arc_summary(content, options)
            if arc_summary:
                body.append("<h3>Арки персонажей</h3>")
                for character_name, points in arc_summary:
                    body.append(f"<h4>{html.escape(character_name)}</h4>")
                    body.append("<ul>")
                    for scene_title, arc_value, description in points:
                        suffix = f": {html.escape(description)}" if description else ""
                        body.append(f"<li><b>{html.escape(scene_title)}</b> - {arc_value:+d}{suffix}</li>")
                    body.append("</ul>")
        for chapter, scenes in self._iter_selected_scenes(content, options):
            chapter_meta = self._chapter_metadata_lines(chapter, options)
            notes = str(chapter.get("notes", "")).strip() if options.include_chapter_notes else ""
            scene_meta: list[tuple[str, list[str]]] = []
            for scene in scenes:
                meta = self._scene_meta_lines(content, scene, options)
                if options.include_character_arcs == "annotations":
                    meta.extend(self._scene_arc_annotations(content, scene))
                if meta:
                    scene_meta.append((self._scene_heading(scene), meta))
            if chapter_meta or notes or scene_meta:
                body.append(f"<h3>{html.escape(chapter.get('title', 'Глава'))} — сведения</h3>")
                if chapter_meta:
                    body.append("<ul>")
                    for item in chapter_meta:
                        label, sep, value = item.partition(": ")
                        if sep:
                            body.append(f"<li><b>{html.escape(label)}:</b> {html.escape(value)}</li>")
                        else:
                            body.append(f"<li>{html.escape(item)}</li>")
                    body.append("</ul>")
                if notes:
                    body.append(f"<p><b>Заметки главы:</b> {html.escape(notes)}</p>")
                for heading, meta in scene_meta:
                    body.append(f"<p><b>{html.escape(heading)}</b></p>")
                    body.append("<ul>")
                    for item in meta:
                        label, sep, value = item.partition(": ")
                        if sep:
                            body.append(f"<li><b>{html.escape(label)}:</b> {html.escape(value)}</li>")
                        else:
                            body.append(f"<li>{html.escape(item)}</li>")
                    body.append("</ul>")
        self._append_library_html(body, content, options)
        return body

    def _additional_materials_docx(self, doc, content: dict, options: ExportOptions) -> None:
        book_meta = self._book_metadata_lines(content.get("book", {}), options)
        if book_meta:
            doc.add_heading("Метаданные книги", level=3)
            for line in book_meta:
                doc.add_paragraph(line)
        if options.include_character_arcs == "summary":
            arc_summary = self._collect_arc_summary(content, options)
            if arc_summary:
                doc.add_heading("Арки персонажей", level=3)
                for character_name, points in arc_summary:
                    doc.add_heading(character_name, level=4)
                    for scene_title, arc_value, description in points:
                        suffix = f": {description}" if description else ""
                        doc.add_paragraph(f"{scene_title} - {arc_value:+d}{suffix}")
        for chapter, scenes in self._iter_selected_scenes(content, options):
            chapter_meta = self._chapter_metadata_lines(chapter, options)
            notes = str(chapter.get("notes", "")).strip() if options.include_chapter_notes else ""
            scene_meta: list[tuple[str, list[str]]] = []
            for scene in scenes:
                meta = self._scene_meta_lines(content, scene, options)
                if options.include_character_arcs == "annotations":
                    meta.extend(self._scene_arc_annotations(content, scene))
                if meta:
                    scene_meta.append((self._scene_heading(scene), meta))
            if chapter_meta or notes or scene_meta:
                doc.add_heading(str(chapter.get("title", "Глава")), level=3)
                for line in chapter_meta:
                    doc.add_paragraph(line)
                if notes:
                    doc.add_paragraph(f"Заметки главы: {notes}")
                for heading, meta in scene_meta:
                    doc.add_heading(heading, level=4)
                    for line in meta:
                        doc.add_paragraph(line)
        self._append_library_docx(doc, content, options)

    def _additional_materials_pdf(self, draw_line, content: dict, options: ExportOptions) -> None:
        book_meta = self._book_metadata_lines(content.get("book", {}), options)
        if book_meta:
            draw_line("Метаданные книги", bold=True)
            for line in book_meta:
                draw_line(line)
            draw_line("")
        if options.include_character_arcs == "summary":
            arc_summary = self._collect_arc_summary(content, options)
            if arc_summary:
                draw_line("Арки персонажей", bold=True)
                for character_name, points in arc_summary:
                    draw_line(character_name, bold=True)
                    for scene_title, arc_value, description in points:
                        suffix = f": {description}" if description else ""
                        draw_line(f"{scene_title} - {arc_value:+d}{suffix}")
                draw_line("")
        for chapter, scenes in self._iter_selected_scenes(content, options):
            chapter_meta = self._chapter_metadata_lines(chapter, options)
            notes = str(chapter.get("notes", "")).strip() if options.include_chapter_notes else ""
            scene_meta: list[tuple[str, list[str]]] = []
            for scene in scenes:
                meta = self._scene_meta_lines(content, scene, options)
                if options.include_character_arcs == "annotations":
                    meta.extend(self._scene_arc_annotations(content, scene))
                if meta:
                    scene_meta.append((self._scene_heading(scene), meta))
            if chapter_meta or notes or scene_meta:
                draw_line(str(chapter.get("title", "Глава")), bold=True)
                for line in chapter_meta:
                    draw_line(line)
                if notes:
                    draw_line(f"Заметки главы: {notes}")
                for heading, meta in scene_meta:
                    draw_line(heading, bold=True)
                    for line in meta:
                        draw_line(line)
                draw_line("")
        self._append_library_pdf(draw_line, content, options)

    def _library_section_enabled(self, options: ExportOptions) -> bool:
        return any(
            (
                options.include_library_characters,
                options.include_library_locations,
                options.include_library_objects,
                options.include_library_plotlines,
            )
        )

    def _append_library_txt(self, lines: list[str], content: dict, options: ExportOptions) -> None:
        if not self._library_section_enabled(options):
            return
        lines.append("")
        lines.append("Библиотека проекта")
        lines.append("-" * 16)
        library = content.get("library", {})
        sections = (
            ("characters", "Персонажи", options.include_library_characters),
            ("locations", "Локации", options.include_library_locations),
            ("objects", "Объекты", options.include_library_objects),
        )
        for kind, title, enabled in sections:
            if not enabled:
                continue
            entities = library.get(kind, [])
            if not entities:
                continue
            lines.append("")
            lines.append(title)
            lines.append("-" * max(5, len(title)))
            for entity in entities:
                if not isinstance(entity, dict):
                    continue
                lines.extend(self._entity_export_lines(content, kind, entity))
                lines.append("")
        if options.include_library_plotlines:
            plotlines = content.get("timeline", {}).get("plotlines", [])
            if plotlines:
                lines.append("")
                lines.append("Сюжетные линии")
                lines.append("-" * 16)
                for plotline in plotlines:
                    if not isinstance(plotline, dict):
                        continue
                    pl_name = str(plotline.get("name", plotline.get("id", "Линия"))).strip()
                    lines.append(pl_name)
                    pl_desc = str(plotline.get("description", "")).strip()
                    if pl_desc:
                        lines.append(f"Описание: {pl_desc}")
                    pl_color = str(plotline.get("color", "")).strip()
                    if pl_color:
                        lines.append(f"Цвет: {pl_color}")
                    lines.append("")

    def _append_library_html(self, body_lines: list[str], content: dict, options: ExportOptions) -> None:
        if not self._library_section_enabled(options):
            return
        body_lines.append("<h3>Библиотека проекта</h3>")
        library = content.get("library", {})
        sections = (
            ("characters", "Персонажи", options.include_library_characters),
            ("locations", "Локации", options.include_library_locations),
            ("objects", "Объекты", options.include_library_objects),
        )
        for kind, title, enabled in sections:
            if not enabled:
                continue
            entities = library.get(kind, [])
            if not entities:
                continue
            body_lines.append(f"<h4>{html.escape(title)}</h4>")
            for entity in entities:
                if not isinstance(entity, dict):
                    continue
                body_lines.append("<div class='library-card'>")
                name = str(entity.get("name", "")).strip() or "Без имени"
                body_lines.append(f"<h5>{html.escape(name)}</h5>")
                for line in self._entity_export_lines(content, kind, entity)[2:]:
                    label, sep, value = line.partition(": ")
                    if sep:
                        body_lines.append(f"<p><b>{html.escape(label)}:</b> {html.escape(value)}</p>")
                    else:
                        body_lines.append(f"<p>{html.escape(line)}</p>")
                body_lines.append("</div>")
        if options.include_library_plotlines:
            plotlines = content.get("timeline", {}).get("plotlines", [])
            if plotlines:
                body_lines.append("<h3>Сюжетные линии</h3>")
                body_lines.append("<ul>")
                for plotline in plotlines:
                    if not isinstance(plotline, dict):
                        continue
                    pl_name = html.escape(str(plotline.get("name", plotline.get("id", "Линия"))))
                    pl_desc = str(plotline.get("description", "")).strip()
                    if pl_desc:
                        body_lines.append(f"<li><b>{pl_name}</b> — {html.escape(pl_desc)}</li>")
                    else:
                        body_lines.append(f"<li>{pl_name}</li>")
                body_lines.append("</ul>")

    def _append_library_docx(self, doc, content: dict, options: ExportOptions) -> None:
        if not self._library_section_enabled(options):
            return
        doc.add_heading("Библиотека проекта", level=3)
        library = content.get("library", {})
        sections = (
            ("characters", "Персонажи", options.include_library_characters),
            ("locations", "Локации", options.include_library_locations),
            ("objects", "Объекты", options.include_library_objects),
        )
        for kind, title, enabled in sections:
            if not enabled:
                continue
            entities = library.get(kind, [])
            if not entities:
                continue
            doc.add_heading(title, level=4)
            for entity in entities:
                if not isinstance(entity, dict):
                    continue
                for line in self._entity_export_lines(content, kind, entity):
                    doc.add_paragraph(line)
                doc.add_paragraph("")
        if options.include_library_plotlines:
            plotlines = content.get("timeline", {}).get("plotlines", [])
            if plotlines:
                doc.add_heading("Сюжетные линии", level=3)
                for plotline in plotlines:
                    if not isinstance(plotline, dict):
                        continue
                    pl_name = str(plotline.get("name", plotline.get("id", "Линия"))).strip()
                    pl_desc = str(plotline.get("description", "")).strip()
                    text = f"{pl_name}: {pl_desc}" if pl_desc else pl_name
                    doc.add_paragraph(text)

    def _append_library_pdf(self, draw_line, content: dict, options: ExportOptions) -> None:
        if not self._library_section_enabled(options):
            return
        draw_line("")
        draw_line("Библиотека проекта", bold=True)
        library = content.get("library", {})
        sections = (
            ("characters", "Персонажи", options.include_library_characters),
            ("locations", "Локации", options.include_library_locations),
            ("objects", "Объекты", options.include_library_objects),
        )
        for kind, title, enabled in sections:
            if not enabled:
                continue
            entities = library.get(kind, [])
            if not entities:
                continue
            draw_line(title, bold=True)
            for entity in entities:
                if not isinstance(entity, dict):
                    continue
                for line in self._entity_export_lines(content, kind, entity):
                    draw_line(line)
                draw_line("")
        if options.include_library_plotlines:
            plotlines = content.get("timeline", {}).get("plotlines", [])
            if plotlines:
                draw_line("Сюжетные линии", bold=True)
                for plotline in plotlines:
                    if not isinstance(plotline, dict):
                        continue
                    pl_name = str(plotline.get("name", plotline.get("id", "Линия"))).strip()
                    pl_desc = str(plotline.get("description", "")).strip()
                    draw_line(f"{pl_name}: {pl_desc}" if pl_desc else pl_name)

    def _entity_export_lines(self, content: dict, kind: str, entity: dict) -> list[str]:
        labels = self._entity_field_labels(kind)
        name = str(entity.get("name", "")).strip() or "Без имени"
        lines = [name, "-" * max(5, len(name))]
        for key in ("extra", "field2", "field3", "short", "arcSummary", "description", "notes"):
            if key == "arcSummary" and kind != "characters":
                continue
            value = str(entity.get(key, "")).strip()
            if value:
                lines.append(f"{labels[key]}: {value}")
        appearances = entity.get("sceneAppearances", [])
        if isinstance(appearances, list) and appearances:
            scene_titles: list[str] = []
            for item in appearances:
                scene_id = item if isinstance(item, str) else str(item.get("sceneId", ""))
                if scene_id:
                    scene_titles.append(self._scene_title_by_id(content, scene_id))
            if scene_titles:
                lines.append("Появления в сценах: " + ", ".join(scene_titles))
        arc_points = entity.get("arcPoints", [])
        if kind == "characters" and isinstance(arc_points, list) and arc_points:
            lines.append("Точки арки:")
            for point in arc_points:
                if not isinstance(point, dict):
                    continue
                scene_id = str(point.get("sceneId", "")).strip()
                scene_title = self._scene_title_by_id(content, scene_id) if scene_id else "—"
                try:
                    arc_value = int(point.get("arcValue", 0))
                except (TypeError, ValueError):
                    arc_value = 0
                description = str(point.get("description", "")).strip()
                suffix = f" — {description}" if description else ""
                lines.append(f"  • {scene_title}: {arc_value:+d}{suffix}")
        return lines

    def _entity_field_labels(self, kind: str) -> dict[str, str]:
        if kind == "characters":
            return {
                "extra": "Роль",
                "field2": "Возраст / этап",
                "field3": "Мотивация",
                "short": "Краткое описание",
                "arcSummary": "Арка персонажа",
                "description": "Подробное описание",
                "notes": "Примечания",
            }
        if kind == "locations":
            return {
                "extra": "Тип локации",
                "field2": "Эпоха / время",
                "field3": "Атмосфера",
                "short": "Краткое описание",
                "description": "Подробное описание",
                "notes": "Примечания",
            }
        return {
            "extra": "Тип объекта",
            "field2": "Состояние",
            "field3": "Назначение",
            "short": "Краткое описание",
            "description": "Подробное описание",
            "notes": "Примечания",
        }

    def _to_txt(self, content: dict, options: ExportOptions) -> str:
        lines: list[str] = []
        book = content.get("book", {})
        lines.append(book.get("title", "Без названия"))
        author = str(book.get("author", "")).strip()
        if author:
            lines.append(f"Автор: {author}")
        lines.append("")

        if options.include_synopsis and book.get("synopsis", "").strip():
            lines.append("Аннотация")
            lines.append(book.get("synopsis", "").strip())
            lines.append("")
        if options.include_toc:
            lines.append("Содержание")
            for volume, entries in self._volume_outline(content, options):
                lines.append(self._volume_label_text(volume))
                for chapter_title, scene_titles in entries:
                    lines.append(f"  - {chapter_title}")
                    for scene_title in scene_titles:
                        lines.append(f"    - {scene_title}")
            lines.append("")

        if options.include_preface and book.get("preface", "").strip():
            lines.append("Предисловие")
            lines.append(book.get("preface", "").strip())
            lines.append("")
        for volume, chapter_blocks in self._volume_blocks(content, options):
            lines.append(self._volume_label_text(volume))
            lines.append("-" * max(5, len(lines[-1])))
            volume_synopsis = str(volume.get("synopsis", "")).strip()
            if options.include_synopsis and volume_synopsis:
                lines.append(volume_synopsis)
                lines.append("")
            for chapter, scenes in chapter_blocks:
                lines.append(chapter.get("title", "Глава"))
                lines.append("-" * max(5, len(lines[-1])))
                chapter_free = self._chapter_free_text_plain(chapter)
                if chapter_free:
                    lines.append(chapter_free)
                    lines.append("")
                for scene in scenes:
                    if str(scene.get("text", "")).strip():
                        lines.append(self._scene_heading(scene))
                        lines.append(scene.get("text", "").strip())
                        lines.append("")
                lines.append("")
            lines.append("")

        if self._has_additional_materials(content, options):
            lines.append("")
            lines.append("Дополнительные материалы")
            lines.append("=" * 24)
            lines.append("")
            lines.extend(self._additional_materials_txt(content, options))
        return "\n".join(lines).strip() + "\n"

    def _to_html(
        self,
        content: dict,
        options: ExportOptions,
        preview_mode: bool = False,
        preview_format: str = "html",
    ) -> str:
        book = content.get("book", {})
        title = html.escape(book.get("title", "Без названия"))
        body_lines: list[str] = []
        if preview_mode:
            body_lines.append(f"<div class='format-badge'>Preview: {html.escape(preview_format.upper())}</div>")
        body_lines.append(f"<h1>{title}</h1>")
        author = str(book.get("author", "")).strip()
        if author:
            body_lines.append(f"<p class='author'><b>Автор:</b> {html.escape(author)}</p>")
        if options.include_synopsis and book.get("synopsis", "").strip():
            body_lines.append("<h2>Аннотация</h2>")
            body_lines.append(f"<p>{html.escape(book.get('synopsis', '')).replace(chr(10), '<br/>')}</p>")
        if options.include_toc:
            body_lines.append("<h2>Содержание</h2>")
            for volume, entries in self._volume_outline(content, options):
                body_lines.append(f"<h3>{html.escape(self._volume_label_text(volume))}</h3>")
                body_lines.append("<ul>")
                for chapter_title, scene_titles in entries:
                    body_lines.append(f"<li><b>{html.escape(chapter_title)}</b></li>")
                    if scene_titles:
                        body_lines.append("<ul>")
                        for scene_title in scene_titles:
                            body_lines.append(f"<li>{html.escape(scene_title)}</li>")
                        body_lines.append("</ul>")
                body_lines.append("</ul>")
        if options.include_preface and book.get("preface", "").strip():
            body_lines.append("<h2>Предисловие</h2>")
            body_lines.append(f"<p>{html.escape(book.get('preface', '')).replace(chr(10), '<br/>')}</p>")
        for volume, chapter_blocks in self._volume_blocks(content, options):
            body_lines.append(f"<h2>{html.escape(self._volume_label_text(volume))}</h2>")
            volume_synopsis = str(volume.get("synopsis", "")).strip()
            if options.include_synopsis and volume_synopsis:
                body_lines.append(f"<p>{html.escape(volume_synopsis).replace(chr(10), '<br/>')}</p>")
            for chapter, scenes in chapter_blocks:
                body_lines.append(f"<h3>{html.escape(chapter.get('title', 'Глава'))}</h3>")
                chapter_free = self._chapter_free_text_html(chapter)
                if chapter_free:
                    body_lines.append(f"<div class='chapter-text'>{chapter_free}</div>")
                for scene in scenes:
                    if not str(scene.get("text", "")).strip():
                        continue
                    body_lines.append(f"<h4>{html.escape(self._scene_heading(scene))}</h4>")
                    text = html.escape(scene.get("text", "")).replace("\n", "<br/>\n")
                    body_lines.append(f"<p>{text}</p>")

        if self._has_additional_materials(content, options):
            body_lines.append("<hr/>")
            body_lines.append("<h2>Дополнительные материалы</h2>")
            body_lines.extend(self._additional_materials_html(content, options))

        css = (
            "body { font-family: '%s', serif; font-size: %dpt; line-height: %.2f; margin: %dpx; color: #1f2937; } "
            "h1, h2, h3 { color: #111827; } ul { margin-top: 0; } "
            ".author { font-size: 12pt; font-style: italic; margin-top: 2px; } "
            ".chapter-text { margin: 8px 0 14px 0; } "
            ".chapter-note { margin: 6px 0 12px 0; color: #0f172a; background: #f8fafc; border-left: 3px solid #64748b; padding: 6px 10px; } "
            ".format-badge { display: inline-block; padding: 4px 8px; border-radius: 999px; "
            "background: #eff6ff; color: #1d4ed8; font-size: 10pt; margin-bottom: 10px; } "
            ".arc-note { color: #6b21a8; background: #faf5ff; border-left: 3px solid #c084fc; padding: 6px 10px; }"
        ) % (options.font_family, max(8, options.font_size), max(1.0, options.line_spacing), 18 if preview_mode else 42)

        return (
            "<!doctype html>\n<html><head><meta charset='utf-8'/>"
            f"<title>{title}</title><style>{css}</style></head><body>\n"
            + "\n".join(body_lines)
            + "\n</body></html>\n"
        )

    def _to_docx(self, content: dict, options: ExportOptions, target: Path) -> Path:
        try:
            from docx import Document
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
            from docx.shared import Pt
        except Exception as err:
            raise ValueError("DOCX export requires python-docx") from err

        doc = Document()
        normal_style = doc.styles["Normal"]
        normal_style.font.name = options.font_family
        normal_style.font.size = Pt(max(8, options.font_size))
        book = content.get("book", {})
        if options.include_header_footer:
            for section in doc.sections:
                section.header.paragraphs[0].text = str(book.get("title", "Scriborium"))
                section.footer.paragraphs[0].text = "Scriborium"
                if options.include_page_numbers:
                    p = section.footer.paragraphs[0]
                    p.add_run(" | Page ")
                    fld = OxmlElement("w:fldSimple")
                    fld.set(qn("w:instr"), "PAGE")
                    p._p.append(fld)
        doc.add_heading(book.get("title", "Без названия"), level=1)
        author = str(book.get("author", "")).strip()
        if author:
            doc.add_paragraph(f"Автор: {author}")
        if options.include_synopsis and book.get("synopsis", "").strip():
            doc.add_heading("Аннотация", level=2)
            doc.add_paragraph(book.get("synopsis", ""))
        if options.include_toc:
            doc.add_heading("Содержание", level=2)
            for volume, entries in self._volume_outline(content, options):
                doc.add_paragraph(self._volume_label_text(volume), style="List Bullet")
                for chapter_title, scene_titles in entries:
                    doc.add_paragraph(chapter_title, style="List Bullet 2")
                    for scene_title in scene_titles:
                        doc.add_paragraph(scene_title, style="List Bullet 3")
        if options.include_preface and book.get("preface", "").strip():
            doc.add_heading("Предисловие", level=2)
            doc.add_paragraph(book.get("preface", ""))

        for volume, chapter_blocks in self._volume_blocks(content, options):
            doc.add_heading(self._volume_label_text(volume), level=2)
            volume_synopsis = str(volume.get("synopsis", "")).strip()
            if options.include_synopsis and volume_synopsis:
                doc.add_paragraph(volume_synopsis)
            for chapter, scenes in chapter_blocks:
                doc.add_heading(chapter.get("title", "Глава"), level=3)
                chapter_free = self._chapter_free_text_plain(chapter)
                if chapter_free:
                    for line in chapter_free.splitlines() or [""]:
                        doc.add_paragraph(line)
                for scene in scenes:
                    if not str(scene.get("text", "")).strip():
                        continue
                    doc.add_heading(self._scene_heading(scene), level=4)
                    for line in scene.get("text", "").splitlines() or [""]:
                        doc.add_paragraph(line)

        if self._has_additional_materials(content, options):
            doc.add_page_break()
            doc.add_heading("Дополнительные материалы", level=2)
            self._additional_materials_docx(doc, content, options)
        doc.save(str(target))
        return target

    def _to_pdf(self, content: dict, options: ExportOptions, target: Path) -> Path:
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.pdfgen import canvas
        except Exception as err:
            raise ValueError("PDF export requires reportlab") from err

        # Unicode font for Cyrillic (DejaVu embeds correctly with reportlab).
        font_name = "Helvetica"
        font_bold = "Helvetica-Bold"
        try:
            import reportlab
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont
            import os as _os
            import sys as _sys

            _pkg_fonts = _os.path.join(
                _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__))),
                "resources",
                "fonts",
            )
            _candidates = [
                _os.path.join(getattr(_sys, "_MEIPASS", ""), "scriborium", "resources", "fonts"),
                _pkg_fonts,
                _os.path.join(_os.path.dirname(reportlab.__file__), "fonts"),
            ]
            _pairs = [("DejaVuSans.ttf", "DejaVuSans-Bold.ttf"), ("Vera.ttf", "VeraBd.ttf")]
            _regular = _bold = ""
            for _fonts_dir in _candidates:
                if not _fonts_dir:
                    continue
                for _reg_name, _bd_name in _pairs:
                    _reg = _os.path.join(_fonts_dir, _reg_name)
                    _bd = _os.path.join(_fonts_dir, _bd_name)
                    if _os.path.exists(_reg) and _os.path.exists(_bd):
                        _regular, _bold = _reg, _bd
                        break
                if _regular and _bold:
                    break
            if _regular and _bold:
                pdfmetrics.registerFont(TTFont("ScrUnicode", _regular))
                pdfmetrics.registerFont(TTFont("ScrUnicode-Bold", _bold))
                font_name = "ScrUnicode"
                font_bold = "ScrUnicode-Bold"
        except Exception:
            pass

        pdf = canvas.Canvas(str(target), pagesize=A4)
        _, height = A4
        x = 40
        y = height - 40
        line_height = max(14, int(options.font_size * max(1.0, options.line_spacing) + 2))
        page_no = 1
        book = content.get("book", {})

        def draw_page_chrome() -> None:
            if not options.include_header_footer:
                return
            pdf.setFont(font_name, 9)
            pdf.drawString(40, height - 20, str(book.get("title", "Scriborium"))[:90])
            footer = "Scriborium"
            if options.include_page_numbers:
                footer = f"{footer} | {page_no}"
            pdf.drawRightString(555, 20, footer)

        def new_page() -> None:
            nonlocal y, page_no
            pdf.showPage()
            page_no += 1
            y = height - 40
            draw_page_chrome()

        max_width = A4[0] - 2 * x  # usable body width

        def draw_line(text: str, bold: bool = False) -> None:
            nonlocal y
            if not text:
                y -= line_height
                return
            size = max(8, options.font_size)
            fname = font_bold if bold else font_name
            pdf.setFont(fname, size)
            first = True
            for part in str(text).split("\n") or [""]:
                chunks = _pdf_wrap(part, max_width, fname, size)
                if not chunks:
                    chunks = [""]
                if not first:
                    chunks = [""] + chunks
                first = False
                for chunk in chunks:
                    if y < 50:
                        new_page()
                        pdf.setFont(fname, size)
                    pdf.drawString(x, y, chunk)
                    y -= line_height

        draw_page_chrome()
        draw_line(book.get("title", "Без названия"), bold=True)
        author = str(book.get("author", "")).strip()
        if author:
            draw_line(f"Автор: {author}")
        draw_line("")
        if options.include_synopsis and book.get("synopsis", "").strip():
            draw_line("Аннотация", bold=True)
            for line in book.get("synopsis", "").splitlines() or [""]:
                draw_line(line)
            draw_line("")
        if options.include_toc:
            draw_line("Содержание", bold=True)
            for volume, entries in self._volume_outline(content, options):
                draw_line(self._volume_label_text(volume), bold=True)
                for chapter_title, scene_titles in entries:
                    draw_line(f"- {chapter_title}")
                    for scene_title in scene_titles:
                        draw_line(f"   - {scene_title}")
            draw_line("")
        if options.include_preface and book.get("preface", "").strip():
            draw_line("Предисловие", bold=True)
            for line in book.get("preface", "").splitlines() or [""]:
                draw_line(line)
            draw_line("")
        for volume, chapter_blocks in self._volume_blocks(content, options):
            draw_line(self._volume_label_text(volume), bold=True)
            volume_synopsis = str(volume.get("synopsis", "")).strip()
            if options.include_synopsis and volume_synopsis:
                for line in volume_synopsis.splitlines() or [""]:
                    draw_line(line)
                draw_line("")
            for chapter, scenes in chapter_blocks:
                draw_line(chapter.get("title", "Глава"), bold=True)
                chapter_free = self._chapter_free_text_plain(chapter)
                if chapter_free:
                    for line in chapter_free.splitlines() or [""]:
                        draw_line(line)
                    draw_line("")
                for scene in scenes:
                    if not str(scene.get("text", "")).strip():
                        continue
                    draw_line(self._scene_heading(scene), bold=True)
                    for line in scene.get("text", "").splitlines() or [""]:
                        draw_line(line)
                    draw_line("")
                draw_line("")
            draw_line("")
        if self._has_additional_materials(content, options):
            draw_line("Дополнительные материалы", bold=True)
            draw_line("")
            self._additional_materials_pdf(draw_line, content, options)
        pdf.save()
        return target


class ExportTemplateService:
    def __init__(self) -> None:
        self.path = app_data_dir() / "export_templates.json"

    def load_all(self) -> list[ExportTemplate]:
        if not self.path.exists():
            return []
        try:
            raw = json.loads(self.path.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            return []

        out: list[ExportTemplate] = []
        for item in raw if isinstance(raw, list) else []:
            options_raw = item.get("options", {})
            out.append(
                ExportTemplate(
                    name=str(item.get("name", "")),
                    fmt=str(item.get("fmt", "txt")),
                    options=ExportOptions(
                        include_preface=bool(options_raw.get("include_preface", True)),
                        include_synopsis=bool(options_raw.get("include_synopsis", True)),
                        selected_scene_ids=list(options_raw.get("selected_scene_ids", [])),
                        include_scene_meta=bool(options_raw.get("include_scene_meta", True)),
                        include_plotlines=bool(options_raw.get("include_plotlines", True)),
                        include_chapter_notes=bool(options_raw.get("include_chapter_notes", False)),
                        include_chapter_metadata=bool(options_raw.get("include_chapter_metadata", True)),
                        include_book_metadata=bool(options_raw.get("include_book_metadata", True)),
                        include_library_characters=bool(options_raw.get("include_library_characters", True)),
                        include_library_locations=bool(options_raw.get("include_library_locations", True)),
                        include_library_objects=bool(options_raw.get("include_library_objects", True)),
                        include_library_plotlines=bool(options_raw.get("include_library_plotlines", True)),
                        include_character_arcs=str(options_raw.get("include_character_arcs", "none")),
                        include_toc=bool(options_raw.get("include_toc", True)),
                        include_header_footer=bool(options_raw.get("include_header_footer", False)),
                        include_page_numbers=bool(options_raw.get("include_page_numbers", False)),
                        font_family=str(options_raw.get("font_family", "Times New Roman")),
                        font_size=int(options_raw.get("font_size", 12)),
                        line_spacing=float(options_raw.get("line_spacing", 1.3)),
                    ),
                )
            )
        return [template for template in out if template.name]

    def save_template(self, template: ExportTemplate) -> None:
        templates = self.load_all()
        filtered = [item for item in templates if item.name != template.name]
        filtered.append(template)
        payload = [
            {
                "name": item.name,
                "fmt": item.fmt,
                "options": asdict(item.options),
            }
            for item in filtered
        ]
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8", newline="\n")










